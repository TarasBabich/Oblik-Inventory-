# -*- coding: utf-8 -*-
"""
Oblik Inventory
===============
Перша робоча desktop-версія програми обліку майна.

Ключові принципи:
- Excel залишається зрозумілим користувачу носієм/експортом.
- pandas використовується для аналізу даних.
- openpyxl використовується для читання/запису Excel зі збереженням структури.
- Flet забезпечує сучасний desktop-інтерфейс.
- зміни НЕ блокуються: програма попереджає і журналює їх;
- "Штатна потреба" не перераховується і переноситься як є;
- "Поточний стан" автоматично перебудовується тільки для поштучного майна,
  яке можна однозначно ідентифікувати за інвентарним або заводським номером.
  Запаси без такого ідентифікатора поки не згортаються автоматично.
"""

from __future__ import annotations

import getpass
import json
import sys
from collections import defaultdict
from copy import copy, deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import flet as ft
import flet_datatable2 as fdt

APP_TITLE = "Oblik Inventory"
APP_VERSION = "0.2.8"

SHEET_STAFF = "Штат"
SHEET_MOVEMENT = "Рух майна"
SHEET_CURRENT = "Поточний стан"
SHEET_SUMMARY = "Зведений"
SHEET_CHANGES = "Контроль змін"
REQUIRED_SHEETS = [SHEET_STAFF, SHEET_MOVEMENT, SHEET_CURRENT, SHEET_SUMMARY, SHEET_CHANGES]
SETTINGS_VIEW = "Налаштування"
TRANSACTION_ID_HEADER = "ID транзакції"
OPERATION_TYPE_HEADER = "Тип операції"

DEFAULT_APP_SETTINGS = {
    "unit_number": "",
    "unit_name": "",
    "commander_position": "Командир військової частини",
    "commander_rank": "",
    "commander_name": "",
    "service_name": "",
    "service_chief_position": "",
    "service_chief_rank": "",
    "service_chief_name": "",
    "finance_chief_position": "Начальник фінансово-економічної служби",
    "finance_chief_rank": "",
    "finance_chief_name": "",
    "index_coefficient_2023": "",
    "index_coefficient_2024": "",
    "index_coefficient_2025": "",
    "document_prefix": "",
    "document_start_number": "1",
    "inventory_prefix": "",
    "inventory_suffix": "",
    "inventory_next_number": "1",
    "inventory_digits": "6",
    "commission_chair_position": "",
    "commission_chair_rank": "",
    "commission_chair_name": "",
    "commission_members": [],
}

MAIN_HEADERS = [
    "№ з/п",
    "Підства \n(номер наряду\n/договір)",
    "Дата підстави \n(наряду\n/договору)",
    "Номер докумету про отримання \n(акт п/п, видаткова накладна)",
    "Дата документу \nпро отримання",
    "Походження \nмайна",
    "Діапазони частот на яких працює РЕБ",
    "Номенклатурний \nкод",
    "Тип майна",
    "Інвентарний номер",
    "Узагальнанена назва згідно номенклатору",
    "Узагальнена назва",
    "Найменування майна згідно документів",
    "Номер матеріалу в SAP",
    "Штатна потреба",
    "Ціна",
    "Кількість",
    "Сума",
    "Заводський \nномер",
    "Окремий підрозді \nбригади де знаходиться на даний момент",
    "Окремий підрозділ \nбатальйону, дивізіону\n де знаходиться на даний момент",
    "Номер документу згідно\n якого підрозділ це \nотримав на даний момент",
    "Дата документу згідно\n якого знаходиться в підрозділі на даний момент",
    "Справний/Несправний \n(переданий в ремонт)/Знищений/Втрачений",
    "Дата якщо пошкоджений/\nвтрачений/знищений",
    "Якщо БПВ або знищений\nномер наказу \nна призначення СР",
    "Дата наказу на \nпризначення СР",
    "Номер наказу \nна результати СР",
    "Дата наказу \nна результати СР",
    "Номер наказу на списання",
    "Дата наказу на списання",
    "Номер Єдиного акту списання",
    "Дата єдиного акту списання",
    "Примітка",
    TRANSACTION_ID_HEADER,
    OPERATION_TYPE_HEADER,
]

SUMMARY_GROUP_FIELDS = {
    "Узагальнена назва": MAIN_HEADERS[11],
    "Узагальнена назва номенклатури": MAIN_HEADERS[10],
    "Номер матеріалу в SAP": MAIN_HEADERS[13],
}

SUMMARY_OUTPUT_HEADERS = [
    "№ з/п",
    "Групування",
    "Значення",
    "Кількість позицій",
    "Загальна кількість",
    "Загальна сума",
    "Справні",
    "Несправні",
    "В ремонті",
    "Знищені",
    "Втрачені",
    "Списані",
]


CHANGE_HEADERS = [
    "№ з/п", "Дата і час", "Тип активу", "Таблиця", "Інв. №",
    "Найменування майна", "Поле", "Було написано", "Стало написано",
    "Тип зміни", "Причина", "Користувач", TRANSACTION_ID_HEADER,
]

STAFF_HEADERS = [
    "№ з/п", "Номенклатурний код", "Найменування згідно номенклатору",
    "Узагальнена кількість за військову частину", "", "",
    "№ з/п", "Номенклатурний код", "Найменування згідно номенклатору",
    "Кількість", "Окремий підрозділ бригади", "Окремий підрозділ батальйону/дивізіону",
]

DUPLICATE_FIELDS = {
    "order_no": MAIN_HEADERS[1],
    "order_date": MAIN_HEADERS[2],
    "receive_doc": MAIN_HEADERS[3],
    "receive_date": MAIN_HEADERS[4],
    "inventory_no": MAIN_HEADERS[9],
    "serial_no": MAIN_HEADERS[18],
}


def is_blank(value: Any) -> bool:
    """Єдина перевірка порожніх значень, включно з pandas NaT/NaN."""
    if value is None:
        return True
    try:
        blank = pd.isna(value)
        return bool(blank)
    except (TypeError, ValueError):
        return False


def norm(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return " ".join(str(value).strip().casefold().split())


def display_value(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    return str(value)


def numeric_value(value: Any) -> Optional[float]:
    """Перетворює число/текст на float для розрахункових колонок."""
    if is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def format_decimal(value: Optional[float], decimals: int = 2) -> str:
    """Людинозрозуміле число для інтерфейсу: 12 500,00."""
    if value is None:
        return "—"
    return f"{value:,.{decimals}f}".replace(",", " ").replace(".", ",")


def calculate_total(price: Any, quantity: Any) -> Optional[float]:
    price_value = numeric_value(price)
    qty_value = numeric_value(quantity)
    if price_value is None or qty_value is None:
        return None
    return price_value * qty_value


def calculate_unit_price(total: Any, quantity: Any) -> Optional[float]:
    total_value = numeric_value(total)
    qty_value = numeric_value(quantity)
    if total_value is None or qty_value is None or qty_value == 0:
        return None
    return total_value / qty_value


def transaction_sequence_width(number: int) -> int:
    """Ширина номера блоками по 3 цифри: 001 ... 999, 001000 ..."""
    digits = len(str(max(1, int(number))))
    return ((digits + 2) // 3) * 3


def format_transaction_id(number: int) -> str:
    number = max(1, int(number))
    width = transaction_sequence_width(number)
    return f"TX-{number:0{width}d}"


def parse_transaction_sequence(value: Any) -> Optional[int]:
    text = str(value or "").strip().upper()
    if not text.startswith("TX-"):
        return None
    numeric = text[3:]
    if not numeric.isdigit():
        return None
    number = int(numeric)
    return number if number > 0 else None


def generate_transaction_id(number: int) -> str:
    """Послідовний стабільний ID облікової операції."""
    return format_transaction_id(number)


def safe_positive_int(value: Any, default: int = 1, maximum: Optional[int] = None) -> int:
    try:
        number = int(str(value).strip())
    except (TypeError, ValueError):
        number = default
    if number < 1:
        number = default
    if maximum is not None:
        number = min(number, maximum)
    return number


def format_inventory_number(prefix: Any, number: int, digits: Any = 6, suffix: Any = "") -> str:
    width = safe_positive_int(digits, 6, 18)
    return f"{str(prefix or '').strip()}{number:0{width}d}{str(suffix or '').strip()}"


def next_inventory_number(settings: dict[str, Any], used_values: list[Any]) -> tuple[str, int]:
    prefix = settings.get("inventory_prefix", "")
    suffix = settings.get("inventory_suffix", "")
    digits = settings.get("inventory_digits", "6")
    number = safe_positive_int(settings.get("inventory_next_number", "1"), 1)
    used = {norm(value) for value in used_values if norm(value)}
    while True:
        candidate = format_inventory_number(prefix, number, digits, suffix)
        if norm(candidate) not in used:
            return candidate, number
        number += 1


def parse_user_value(header: str, text: str) -> Any:
    text = text.strip()
    if not text:
        return None
    h = header.casefold()
    if "дата" in h:
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
        return text
    if header in ("Ціна", "Кількість", "Сума", "Штатна потреба"):
        candidate = text.replace(" ", "").replace(",", ".")
        try:
            return float(candidate)
        except ValueError:
            return text
    return text


@dataclass
class DuplicateInfo:
    score: int = 0
    other_excel_rows: tuple[int, ...] = ()


class OblikWorkbook:
    """Excel-книга + основна логіка обліку."""

    def __init__(self) -> None:
        self.path: Optional[Path] = None
        self.wb = None
        self.dirty = False

    def create_new(self, path: Path) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        self._create_main_sheet(wb, SHEET_MOVEMENT)
        self._create_main_sheet(wb, SHEET_CURRENT)

        ws_staff = wb.create_sheet(SHEET_STAFF)
        for col, header in enumerate(STAFF_HEADERS, 1):
            ws_staff.cell(1, col, header)
        self._style_header(ws_staff, 1, len(STAFF_HEADERS))
        ws_staff.freeze_panes = "A2"

        ws_summary = wb.create_sheet(SHEET_SUMMARY)
        for col, header in enumerate(SUMMARY_OUTPUT_HEADERS, 1):
            ws_summary.cell(1, col, header)
        self._style_header(ws_summary, 1, len(SUMMARY_OUTPUT_HEADERS))
        ws_summary.freeze_panes = "A2"

        ws_changes = wb.create_sheet(SHEET_CHANGES)
        for col, header in enumerate(CHANGE_HEADERS, 1):
            ws_changes.cell(1, col, header)
        self._style_header(ws_changes, 1, len(CHANGE_HEADERS))
        ws_changes.freeze_panes = "A2"

        wb.save(path)
        self.load(path)

    def _create_main_sheet(self, wb: Workbook, name: str) -> None:
        ws = wb.create_sheet(name)
        for col, header in enumerate(MAIN_HEADERS, 1):
            ws.cell(1, col, header)
        self._style_header(ws, 1, len(MAIN_HEADERS))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(MAIN_HEADERS))}1"
        for idx in range(1, len(MAIN_HEADERS) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 18

    def _style_header(self, ws, row: int, count: int) -> None:
        fill = PatternFill("solid", fgColor="1F4E78")
        side = Side(style="thin", color="A0A0A0")
        for col in range(1, count + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=side, right=side, top=side, bottom=side)
        ws.row_dimensions[row].height = 60

    def load(self, path: Path) -> None:
        self.wb = load_workbook(path)
        missing = [s for s in REQUIRED_SHEETS if s not in self.wb.sheetnames]
        if missing:
            raise ValueError("У книзі відсутні аркуші: " + ", ".join(missing))

        # Міграція старого шаблону: раніше рядок 2 містив службові числа 1..34.
        # Тепер дані починаються з рядка 2, щоб № з/п міг бути =ROW()-1.
        migrated = False
        for sheet_name in (SHEET_MOVEMENT, SHEET_CURRENT):
            ws = self.wb[sheet_name]
            expected = list(range(1, min(ws.max_column, len(MAIN_HEADERS)) + 1))
            actual = [ws.cell(2, c).value for c in range(1, len(expected) + 1)]
            normalized = []
            for value in actual:
                try:
                    normalized.append(int(float(value)))
                except (TypeError, ValueError):
                    normalized.append(None)
            if normalized == expected:
                ws.delete_rows(2, 1)
                migrated = True

            if self._ensure_main_sheet_schema(ws):
                migrated = True

            # Після міграції/відкриття приводимо розрахункові колонки до правил.
            self._restore_row_formulas(ws)

        if self._ensure_change_sheet_schema(self.wb[SHEET_CHANGES]):
            migrated = True
        if self._ensure_transaction_ids():
            migrated = True
        if self._ensure_operation_types():
            migrated = True

        self.path = path
        self.dirty = migrated

    def _ensure_main_sheet_schema(self, ws) -> bool:
        """Додає нові технічні колонки в кінець, не зсуваючи стару Excel-структуру."""
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        changed = False
        for technical_header in (TRANSACTION_ID_HEADER, OPERATION_TYPE_HEADER):
            if technical_header in headers:
                continue
            col = ws.max_column + 1
            ws.cell(1, col, technical_header)
            if col > 1 and ws.cell(1, col - 1).has_style:
                ws.cell(1, col)._style = copy(ws.cell(1, col - 1)._style)
            ws.column_dimensions[get_column_letter(col)].width = 18
            headers.append(technical_header)
            changed = True
        if changed:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"
        return changed

    def _ensure_change_sheet_schema(self, ws) -> bool:
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        if TRANSACTION_ID_HEADER in headers:
            return False
        col = ws.max_column + 1
        ws.cell(1, col, TRANSACTION_ID_HEADER)
        if col > 1 and ws.cell(1, col - 1).has_style:
            ws.cell(1, col)._style = copy(ws.cell(1, col - 1)._style)
        ws.column_dimensions[get_column_letter(col)].width = 18
        return True

    def _next_transaction_sequence(self) -> int:
        """Наступний номер з урахуванням руху й аудиту, тому видалені ID не повторюються."""
        largest = 0
        for sheet_name in (SHEET_MOVEMENT, SHEET_CHANGES):
            ws = self.wb[sheet_name]
            headers = self.headers(sheet_name)
            if TRANSACTION_ID_HEADER not in headers:
                continue
            tx_col = headers.index(TRANSACTION_ID_HEADER) + 1
            for row in range(2, ws.max_row + 1):
                sequence = parse_transaction_sequence(ws.cell(row, tx_col).value)
                if sequence is not None:
                    largest = max(largest, sequence)
        return largest + 1

    def _replace_transaction_id_in_sheet(
        self,
        sheet_name: str,
        old_id: str,
        new_id: str,
    ) -> bool:
        ws = self.wb[sheet_name]
        headers = self.headers(sheet_name)
        if TRANSACTION_ID_HEADER not in headers:
            return False
        tx_col = headers.index(TRANSACTION_ID_HEADER) + 1
        changed = False
        wanted = norm(old_id)
        for row in range(2, ws.max_row + 1):
            if norm(ws.cell(row, tx_col).value) == wanted:
                ws.cell(row, tx_col, new_id)
                changed = True
        return changed

    def _ensure_transaction_ids(self) -> bool:
        """Мігрує UUID/порожні ID у послідовні TX-001, TX-002 ... без повторного перенумерування."""
        ws = self.wb[SHEET_MOVEMENT]
        headers = self.headers(SHEET_MOVEMENT)
        if TRANSACTION_ID_HEADER not in headers:
            return False
        tx_col = headers.index(TRANSACTION_ID_HEADER) + 1
        changed = False
        seen_sequences: set[int] = set()

        # Спочатку визначаємо найбільший уже чинний послідовний номер.
        next_sequence = self._next_transaction_sequence()

        for row in range(2, ws.max_row + 1):
            if not self._row_has_data(ws, row, 2, tx_col - 1):
                continue

            current = str(ws.cell(row, tx_col).value or "").strip()
            sequence = parse_transaction_sequence(current)
            if sequence is not None and sequence not in seen_sequences:
                seen_sequences.add(sequence)
                continue

            new_id = format_transaction_id(next_sequence)
            while next_sequence in seen_sequences or self._transaction_id_exists(
                new_id, ignore_excel_row=row
            ):
                next_sequence += 1
                new_id = format_transaction_id(next_sequence)

            old_id = current
            ws.cell(row, tx_col, new_id)
            seen_sequences.add(next_sequence)
            next_sequence += 1
            changed = True

            # Для старих UUID синхронно оновлюємо посилання в журналі та поточному стані.
            if old_id and parse_transaction_sequence(old_id) is None:
                self._replace_transaction_id_in_sheet(SHEET_CHANGES, old_id, new_id)
                self._replace_transaction_id_in_sheet(SHEET_CURRENT, old_id, new_id)

        return changed

    def _ensure_operation_types(self) -> bool:
        ws = self.wb[SHEET_MOVEMENT]
        headers = self.headers(SHEET_MOVEMENT)
        if OPERATION_TYPE_HEADER not in headers:
            return False
        op_col = headers.index(OPERATION_TYPE_HEADER) + 1
        tx_col = headers.index(TRANSACTION_ID_HEADER) + 1
        changed = False
        for row in range(2, ws.max_row + 1):
            if not self._row_has_data(ws, row, 2, tx_col):
                continue
            if not norm(ws.cell(row, op_col).value):
                ws.cell(row, op_col, "Імпортований запис")
                changed = True
        return changed

    def _transaction_id_exists(self, transaction_id: Any, ignore_excel_row: Optional[int] = None) -> bool:
        wanted = norm(transaction_id)
        if not wanted or self.wb is None:
            return False
        df = self.dataframe(SHEET_MOVEMENT)
        if df.empty or TRANSACTION_ID_HEADER not in df.columns:
            return False
        for _, row in df.iterrows():
            if ignore_excel_row and int(row["_excel_row"]) == ignore_excel_row:
                continue
            if norm(row.get(TRANSACTION_ID_HEADER)) == wanted:
                return True
        return False

    def save(self, path: Optional[Path] = None) -> None:
        if self.wb is None:
            raise RuntimeError("Книгу не відкрито")
        target = path or self.path
        if not target:
            raise RuntimeError("Не вказано шлях для збереження")
        self.wb.save(target)
        self.path = Path(target)
        self.dirty = False

    def headers(self, sheet_name: str) -> list[str]:
        ws = self.wb[sheet_name]
        return [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]

    def _row_has_data(self, ws, row: int, start_col: int = 1, end_col: Optional[int] = None) -> bool:
        end_col = end_col or ws.max_column
        return any(ws.cell(row, c).value not in (None, "") for c in range(start_col, end_col + 1))

    def _first_empty_data_row(self, ws, start_row: int = 2) -> int:
        upper = max(ws.max_row + 2, start_row + 2)
        for row in range(start_row, upper + 1):
            if not self._row_has_data(ws, row, 2, ws.max_column):
                return row
        return upper + 1

    def _restore_row_formulas(self, ws) -> None:
        """Відновлює формули нумерації та суми для всіх непорожніх рядків."""
        for row in range(2, ws.max_row + 1):
            if not self._row_has_data(ws, row, 2, ws.max_column):
                continue
            ws.cell(row, 1, f"=ROW()-1")
            ws.cell(row, 18, f"=P{row}*Q{row}")

    def dataframe(self, sheet_name: str) -> pd.DataFrame:
        ws = self.wb[sheet_name]
        headers = self.headers(sheet_name)
        rows = []
        start_row = 2
        for row_no in range(start_row, ws.max_row + 1):
            values = [ws.cell(row_no, col).value for col in range(1, len(headers) + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            row = {headers[i]: values[i] for i in range(len(headers))}
            row["_excel_row"] = row_no
            rows.append(row)
        return pd.DataFrame(rows)

    def append_record(self, sheet_name: str, values: dict[str, Any], reason: str = "") -> int:
        ws = self.wb[sheet_name]
        headers = self.headers(sheet_name)
        values = dict(values)
        if sheet_name == SHEET_MOVEMENT:
            sequence = self._next_transaction_sequence()
            transaction_id = format_transaction_id(sequence)
            while self._transaction_id_exists(transaction_id):
                sequence += 1
                transaction_id = format_transaction_id(sequence)
            values[TRANSACTION_ID_HEADER] = transaction_id
            if not norm(values.get(OPERATION_TYPE_HEADER)):
                values[OPERATION_TYPE_HEADER] = "Первинний запис"
        start_row = 2
        row = self._first_empty_data_row(ws, start_row)
        template_row = start_row

        for col in range(1, len(headers) + 1):
            if template_row <= ws.max_row:
                src = ws.cell(template_row, col)
                dst = ws.cell(row, col)
                if src.has_style:
                    dst._style = copy(src._style)

        for col, header in enumerate(headers, 1):
            if header == "№ з/п":
                ws.cell(row, col, "=ROW()-1")
            elif header == "Сума":
                ws.cell(row, col, f"=P{row}*Q{row}")
            else:
                ws.cell(row, col, values.get(header))

        self.dirty = True
        if sheet_name == SHEET_MOVEMENT:
            self._log_change(values, sheet_name, "Весь запис", "", self._record_summary(values), "Додавання", reason)
        return row

    def update_record(self, sheet_name: str, excel_row: int, new_values: dict[str, Any], reason: str = ""):
        ws = self.wb[sheet_name]
        headers = self.headers(sheet_name)
        old_values = {header: ws.cell(excel_row, i + 1).value for i, header in enumerate(headers)}
        changes = []

        for i, header in enumerate(headers, 1):
            if header in ("№ з/п", "Сума", TRANSACTION_ID_HEADER, OPERATION_TYPE_HEADER):
                continue
            old = old_values.get(header)
            new = new_values.get(header)
            if norm(old) != norm(new):
                ws.cell(excel_row, i, new)
                changes.append((header, old, new))

        # Розрахункові колонки ніколи не вводяться вручну.
        ws.cell(excel_row, 1, "=ROW()-1")
        ws.cell(excel_row, 18, f"=P{excel_row}*Q{excel_row}")

        if changes:
            self.dirty = True
            merged = dict(old_values)
            merged.update(new_values)
            merged[TRANSACTION_ID_HEADER] = old_values.get(TRANSACTION_ID_HEADER, "")
            for field, old, new in changes:
                self._log_change(merged, sheet_name, field, old, new, "Виправлення", reason)
        return changes

    def delete_record(self, sheet_name: str, excel_row: int, reason: str = "") -> None:
        ws = self.wb[sheet_name]
        headers = self.headers(sheet_name)
        old_values = {header: ws.cell(excel_row, i + 1).value for i, header in enumerate(headers)}
        summary = self._record_summary(old_values)
        for col in range(1, len(headers) + 1):
            ws.cell(excel_row, col).value = None
        self.dirty = True
        self._log_change(old_values, sheet_name, "Весь запис", summary, "", "Видалення", reason)

    def _record_summary(self, values: dict[str, Any]) -> str:
        keys = [TRANSACTION_ID_HEADER, MAIN_HEADERS[9], MAIN_HEADERS[18], MAIN_HEADERS[12], MAIN_HEADERS[16], MAIN_HEADERS[19], MAIN_HEADERS[20]]
        parts = []
        for key in keys:
            value = values.get(key)
            if value not in (None, ""):
                parts.append(f"{key.replace(chr(10), ' ')}={display_value(value)}")
        return "; ".join(parts)[:1000]

    def _log_change(self, record, table, field, old, new, change_type, reason):
        ws = self.wb[SHEET_CHANGES]
        row = 2
        while row <= max(ws.max_row + 1, 2):
            if not any(
                ws.cell(row, col).value not in (None, "")
                for col in range(2, len(CHANGE_HEADERS) + 1)
            ):
                break
            row += 1

        seq = 1
        for r in range(2, row):
            try:
                seq = max(seq, int(ws.cell(r, 1).value or 0) + 1)
            except (TypeError, ValueError):
                pass

        vals = [
            seq, datetime.now(), record.get(MAIN_HEADERS[8], ""), table,
            record.get(MAIN_HEADERS[9], ""),
            record.get(MAIN_HEADERS[12], "") or record.get(MAIN_HEADERS[11], ""),
            field, display_value(old), display_value(new), change_type,
            reason, getpass.getuser(), record.get(TRANSACTION_ID_HEADER, ""),
        ]
        for col, value in enumerate(vals, 1):
            ws.cell(row, col, value)
        ws.cell(row, 2).number_format = "dd.mm.yyyy hh:mm:ss"

    def duplicate_map(self) -> dict[int, DuplicateInfo]:
        df = self.dataframe(SHEET_MOVEMENT)
        if df.empty:
            return {}

        records = []
        for _, row in df.iterrows():
            operation_type = norm(row.get(OPERATION_TYPE_HEADER))
            if operation_type and operation_type not in ("первинний запис", "імпортований запис"):
                continue
            identifier = norm(row.get(DUPLICATE_FIELDS["inventory_no"])) or norm(row.get(DUPLICATE_FIELDS["serial_no"]))
            sig = [
                norm(row.get(DUPLICATE_FIELDS["order_no"])),
                norm(row.get(DUPLICATE_FIELDS["order_date"])),
                norm(row.get(DUPLICATE_FIELDS["receive_doc"])),
                norm(row.get(DUPLICATE_FIELDS["receive_date"])),
                identifier,
            ]
            records.append((int(row["_excel_row"]), sig))

        inverted = [defaultdict(set) for _ in range(5)]
        for excel_row, sig in records:
            for i, value in enumerate(sig):
                if value:
                    inverted[i][value].add(excel_row)

        result = {}
        for excel_row, sig in records:
            counts = defaultdict(int)
            for i, value in enumerate(sig):
                if not value:
                    continue
                for other in inverted[i].get(value, set()):
                    if other != excel_row:
                        counts[other] += 1
            if counts:
                max_score = max(counts.values())
                others = tuple(sorted(k for k, score in counts.items() if score == max_score))
                result[excel_row] = DuplicateInfo(max_score, others)
            else:
                result[excel_row] = DuplicateInfo()
        return result

    def score_candidate_duplicate(self, candidate, ignore_excel_row=None) -> DuplicateInfo:
        df = self.dataframe(SHEET_MOVEMENT)
        if df.empty:
            return DuplicateInfo()

        candidate_sig = [
            norm(candidate.get(DUPLICATE_FIELDS["order_no"])),
            norm(candidate.get(DUPLICATE_FIELDS["order_date"])),
            norm(candidate.get(DUPLICATE_FIELDS["receive_doc"])),
            norm(candidate.get(DUPLICATE_FIELDS["receive_date"])),
            norm(candidate.get(DUPLICATE_FIELDS["inventory_no"])) or norm(candidate.get(DUPLICATE_FIELDS["serial_no"])),
        ]

        best = 0
        rows = []
        for _, row in df.iterrows():
            rnum = int(row["_excel_row"])
            if ignore_excel_row and rnum == ignore_excel_row:
                continue
            operation_type = norm(row.get(OPERATION_TYPE_HEADER))
            if operation_type and operation_type not in ("первинний запис", "імпортований запис"):
                continue
            sig = [
                norm(row.get(DUPLICATE_FIELDS["order_no"])),
                norm(row.get(DUPLICATE_FIELDS["order_date"])),
                norm(row.get(DUPLICATE_FIELDS["receive_doc"])),
                norm(row.get(DUPLICATE_FIELDS["receive_date"])),
                norm(row.get(DUPLICATE_FIELDS["inventory_no"])) or norm(row.get(DUPLICATE_FIELDS["serial_no"])),
            ]
            score = sum(1 for a, b in zip(candidate_sig, sig) if a and b and a == b)
            if score > best:
                best = score
                rows = [rnum]
            elif score == best and score > 0:
                rows.append(rnum)
        return DuplicateInfo(best, tuple(rows))

    def build_summary_dataframe(self, group_header: str) -> pd.DataFrame:
        """Агрегує актуальний Поточний стан за вибраним полем."""
        if group_header not in SUMMARY_GROUP_FIELDS.values():
            raise ValueError(f"Непідтримуване поле групування: {group_header}")

        current = self.dataframe(SHEET_CURRENT)
        if current.empty:
            return pd.DataFrame(columns=SUMMARY_OUTPUT_HEADERS)

        groups: dict[str, dict[str, Any]] = {}
        for _, row in current.iterrows():
            raw_group = display_value(row.get(group_header)).strip()
            group_value = raw_group or "Не вказано"

            quantity = numeric_value(row.get(MAIN_HEADERS[16]))
            if quantity is None:
                # Для поштучного майна без явно заданої кількості один
                # поточний рядок означає одну одиницю.
                quantity = 1.0 if (
                    norm(row.get(MAIN_HEADERS[9]))
                    or norm(row.get(MAIN_HEADERS[18]))
                ) else 0.0

            price = numeric_value(row.get(MAIN_HEADERS[15]))
            total = (price * quantity) if price is not None else 0.0
            status = norm(row.get(MAIN_HEADERS[23]))

            bucket = groups.setdefault(
                group_value,
                {
                    "Кількість позицій": 0,
                    "Загальна кількість": 0.0,
                    "Загальна сума": 0.0,
                    "Справні": 0.0,
                    "Несправні": 0.0,
                    "В ремонті": 0.0,
                    "Знищені": 0.0,
                    "Втрачені": 0.0,
                    "Списані": 0.0,
                },
            )
            bucket["Кількість позицій"] += 1
            bucket["Загальна кількість"] += quantity
            bucket["Загальна сума"] += total

            if "передан" in status and "ремонт" in status:
                bucket["В ремонті"] += quantity
            elif "несправ" in status:
                bucket["Несправні"] += quantity
            elif "знищ" in status:
                bucket["Знищені"] += quantity
            elif "втрач" in status:
                bucket["Втрачені"] += quantity
            elif "спис" in status:
                bucket["Списані"] += quantity
            elif "справ" in status:
                bucket["Справні"] += quantity

        rows = []
        group_label = next(
            label for label, header in SUMMARY_GROUP_FIELDS.items()
            if header == group_header
        )
        for index, group_value in enumerate(
            sorted(groups, key=lambda value: value.casefold()),
            start=1,
        ):
            bucket = groups[group_value]
            rows.append({
                "№ з/п": index,
                "Групування": group_label,
                "Значення": group_value,
                **bucket,
            })

        return pd.DataFrame(rows, columns=SUMMARY_OUTPUT_HEADERS)

    def rebuild_summary(self, group_header: str) -> pd.DataFrame:
        """Перебудовує аркуш Зведений відповідно до вибраного режиму."""
        summary = self.build_summary_dataframe(group_header)
        ws = self.wb[SHEET_SUMMARY]

        # Зведений є похідним аркушем, тому його вміст можна безпечно
        # перебудовувати з Поточного стану.
        if ws.max_row > 0:
            ws.delete_rows(1, ws.max_row)

        for col, header in enumerate(SUMMARY_OUTPUT_HEADERS, 1):
            ws.cell(1, col, header)
        self._style_header(ws, 1, len(SUMMARY_OUTPUT_HEADERS))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(SUMMARY_OUTPUT_HEADERS))}1"

        for row_index, record in enumerate(summary.to_dict("records"), start=2):
            for col_index, header in enumerate(SUMMARY_OUTPUT_HEADERS, start=1):
                ws.cell(row_index, col_index, record.get(header))

        widths = {
            "A": 8,
            "B": 32,
            "C": 42,
            "D": 18,
            "E": 18,
            "F": 20,
            "G": 14,
            "H": 14,
            "I": 14,
            "J": 14,
            "K": 14,
            "L": 14,
        }
        for column, width in widths.items():
            ws.column_dimensions[column].width = width

        self.dirty = True
        return summary

    def rebuild_current_state(self) -> tuple[int, int]:
        """Поточний стан по поштучному майну. Штатну потребу не перераховує."""
        df = self.dataframe(SHEET_MOVEMENT)
        ws_cur = self.wb[SHEET_CURRENT]
        if df.empty:
            return 0, 0

        headers = self.headers(SHEET_MOVEMENT)
        name_col = MAIN_HEADERS[12]
        inv_col = MAIN_HEADERS[9]
        serial_col = MAIN_HEADERS[18]
        event_date_cols = [
            MAIN_HEADERS[22],  # документ поточного розміщення
            MAIN_HEADERS[24],  # дата зміни стану
            MAIN_HEADERS[30],  # дата наказу на списання
            MAIN_HEADERS[32],  # дата єдиного акту списання
            MAIN_HEADERS[4],   # дата первинного отримання
        ]

        def key_for(row) -> str:
            inv = norm(row.get(inv_col))
            if inv:
                return "INV|" + inv
            serial = norm(row.get(serial_col))
            if serial:
                return "SER|" + serial + "|" + norm(row.get(name_col))
            return ""

        df = df.copy()
        df["_asset_key"] = df.apply(key_for, axis=1)
        identified = df[df["_asset_key"] != ""].copy()
        skipped = int((df["_asset_key"] == "").sum())

        def as_ts(value):
            if value is None or value == "":
                return pd.NaT
            return pd.to_datetime(value, dayfirst=True, errors="coerce")

        event_dates = pd.DataFrame(
            {
                header: identified[header].map(as_ts)
                for header in event_date_cols
                if header in identified.columns
            },
            index=identified.index,
        )
        if event_dates.empty:
            identified["_sort_date"] = pd.Timestamp("1900-01-01")
        else:
            identified["_sort_date"] = event_dates.max(axis=1)
            identified["_sort_date"] = identified["_sort_date"].fillna(
                pd.Timestamp("1900-01-01")
            )
        identified = identified.sort_values(["_asset_key", "_sort_date", "_excel_row"])
        latest = identified.groupby("_asset_key", as_index=False).tail(1)

        # Кількісні позиції/запаси без унікального номера поки не
        # згортаємо автоматично. Зберігаємо їхні поточні рядки без змін,
        # щоб перебудова поштучного майна не призвела до втрати даних.
        existing_current = self.dataframe(SHEET_CURRENT)
        preserved_rows = []
        if not existing_current.empty:
            for _, current_row in existing_current.iterrows():
                inv = norm(current_row.get(inv_col))
                serial = norm(current_row.get(serial_col))
                if not inv and not serial:
                    preserved_rows.append({h: current_row.get(h) for h in headers})

        for row in range(2, ws_cur.max_row + 1):
            for col in range(1, min(len(headers), ws_cur.max_column) + 1):
                ws_cur.cell(row, col).value = None

        out_row = 2
        for _, record in latest.iterrows():
            for col, header in enumerate(headers, 1):
                if header == "№ з/п":
                    value = "=ROW()-1"
                elif header == "Сума":
                    value = f"=P{out_row}*Q{out_row}"
                else:
                    value = record.get(header)
                ws_cur.cell(out_row, col, value)
            out_row += 1

        for record in preserved_rows:
            for col, header in enumerate(headers, 1):
                if header == "№ з/п":
                    value = "=ROW()-1"
                elif header == "Сума":
                    value = f"=P{out_row}*Q{out_row}"
                else:
                    value = record.get(header)
                ws_cur.cell(out_row, col, value)
            out_row += 1

        self.dirty = True
        return len(latest), skipped



MAX_TABLE_ROWS = 250


class AppSettingsStore:
    """Локальні portable-налаштування програми у JSON поруч із Oblik.exe."""

    def __init__(self) -> None:
        base = (
            Path(sys.executable).parent
            if getattr(sys, "frozen", False)
            else Path(__file__).resolve().parent.parent
        )
        self.path = base / "oblik_settings.json"

    def load(self) -> dict[str, Any]:
        values = deepcopy(DEFAULT_APP_SETTINGS)
        if not self.path.exists():
            return values
        try:
            loaded = json.loads(self.path.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                for key in values:
                    if key == "commission_members":
                        members = loaded.get(key, [])
                        if isinstance(members, list):
                            cleaned = []
                            for member in members:
                                if not isinstance(member, dict):
                                    continue
                                cleaned.append({
                                    "position": str(member.get("position", "") or "").strip(),
                                    "rank": str(member.get("rank", "") or "").strip(),
                                    "name": str(member.get("name", "") or "").strip(),
                                })
                            values[key] = cleaned
                    elif key in loaded and loaded[key] is not None:
                        values[key] = str(loaded[key])
        except (OSError, json.JSONDecodeError):
            # Пошкоджені налаштування не повинні блокувати запуск програми.
            pass
        return values

    def save(self, values: dict[str, Any]) -> Path:
        data = deepcopy(DEFAULT_APP_SETTINGS)
        for key in data:
            if key == "commission_members":
                members = values.get(key, [])
                cleaned = []
                if isinstance(members, list):
                    for member in members:
                        if not isinstance(member, dict):
                            continue
                        cleaned.append({
                            "position": str(member.get("position", "") or "").strip(),
                            "rank": str(member.get("rank", "") or "").strip(),
                            "name": str(member.get("name", "") or "").strip(),
                        })
                data[key] = cleaned
            else:
                value = values.get(key, "")
                data[key] = "" if value is None else str(value).strip()
        self.path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return self.path


class FletOblikApp:
    """Flet-інтерфейс. Облікова логіка лишається в OblikWorkbook."""

    def __init__(self, page: ft.Page):
        self.page = page
        self.model = OblikWorkbook()
        self.current_sheet = SHEET_MOVEMENT
        self.selected_excel_row: Optional[int] = None
        self.file_picker = ft.FilePicker()
        self.settings_store = AppSettingsStore()
        self.settings = self.settings_store.load()
        self.settings_fields: dict[str, ft.TextField] = {}
        self.commission_member_entries: list[dict[str, Any]] = []
        self.summary_group_mode = "Узагальнена назва"

        self.page.title = f"{APP_TITLE} {APP_VERSION}"
        self.page.theme_mode = ft.ThemeMode.LIGHT
        self.page.theme = ft.Theme(color_scheme_seed=ft.Colors.BLUE)
        self.page.padding = 0
        self.page.spacing = 0
        self.page.window.width = 1500
        self.page.window.height = 900
        self.page.window.min_width = 1000
        self.page.window.min_height = 650

        self.status = ft.Text(
            "Відкрийте існуючу книгу або створіть нову",
            size=12,
            color=ft.Colors.BLUE_GREY_700,
        )
        self.file_label = ft.Text("Файл не відкрито", weight=ft.FontWeight.W_600)
        self.search = ft.TextField(
            hint_text="Пошук у відкритій таблиці…",
            prefix_icon=ft.Icons.SEARCH,
            on_change=self._search_changed,
            expand=True,
        )
        # Таблиця сама керує вертикальною/горизонтальною прокруткою через DataTable2.
        self.table_host = ft.Column(expand=True)

        self.btn_add = ft.Button(
            content="+ Додати запис",
            icon=ft.Icons.ADD,
            on_click=self._add_record,
            disabled=True,
        )
        self.btn_edit = ft.Button(
            content="Редагувати",
            icon=ft.Icons.EDIT,
            on_click=self._edit_record,
            disabled=True,
        )
        self.btn_delete = ft.Button(
            content="Видалити",
            icon=ft.Icons.DELETE_OUTLINE,
            on_click=self._delete_record,
            disabled=True,
        )

        self.summary_group_dropdown = ft.Dropdown(
            label="Групувати за",
            value=self.summary_group_mode,
            width=320,
            visible=False,
            options=[
                ft.DropdownOption(key=label, text=label)
                for label in SUMMARY_GROUP_FIELDS
            ],
            on_select=self._summary_group_changed,
        )
        self.btn_summary_refresh = ft.Button(
            content="Оновити зведений",
            icon=ft.Icons.REFRESH,
            visible=False,
            on_click=self._refresh_summary,
        )

        self.action_menu = ft.PopupMenuButton(
            disabled=True,
            menu_position=ft.PopupMenuPosition.UNDER,
            content=ft.Container(
                padding=ft.Padding.symmetric(horizontal=14, vertical=9),
                border=ft.Border.all(1, ft.Colors.BLUE_200),
                border_radius=22,
                bgcolor=ft.Colors.BLUE_50,
                content=ft.Row(
                    tight=True,
                    spacing=6,
                    controls=[
                        ft.Icon(ft.Icons.BOLT, size=18, color=ft.Colors.BLUE_700),
                        ft.Text("Дія", weight=ft.FontWeight.W_600, color=ft.Colors.BLUE_800),
                        ft.Icon(ft.Icons.ARROW_DROP_DOWN, size=20, color=ft.Colors.BLUE_700),
                    ],
                ),
            ),
            items=[
                ft.PopupMenuItem(
                    icon=ft.Icons.SWAP_HORIZ,
                    content="Перемістити",
                    on_click=self._open_move_dialog,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.INVENTORY_2,
                    content="Видати запас",
                    on_click=self._open_issue_stock_dialog,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.ASSIGNMENT_RETURN,
                    content="Повернути запас",
                    on_click=self._open_return_stock_dialog,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.DELETE_SWEEP,
                    content="Списати",
                    on_click=self._open_writeoff_dialog,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.HEALTH_AND_SAFETY,
                    content="Змінити стан",
                    on_click=self._open_status_dialog,
                ),
                ft.PopupMenuItem(content=ft.Divider(height=1), height=12),
                ft.PopupMenuItem(
                    icon=ft.Icons.HISTORY,
                    content="Переглянути історію",
                    on_click=self._show_selected_history,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.EDIT,
                    content="Редагувати запис",
                    on_click=self._edit_record,
                ),
                ft.PopupMenuItem(
                    icon=ft.Icons.DELETE_OUTLINE,
                    content="Видалити запис",
                    on_click=self._delete_record,
                ),
            ],
        )

        self.sheet_buttons: dict[str, ft.Button] = {}
        self._build_page()
        self._autoload_default_book()

    def _build_page(self):
        toolbar = ft.Container(
            bgcolor=ft.Colors.BLUE_GREY_900,
            padding=ft.Padding.symmetric(horizontal=14, vertical=10),
            content=ft.Row(
                controls=[
                    ft.Text(
                        "OBLIK",
                        size=20,
                        weight=ft.FontWeight.BOLD,
                        color=ft.Colors.WHITE,
                    ),
                    ft.VerticalDivider(color=ft.Colors.WHITE_24),
                    ft.Button(
                        content="Відкрити Excel",
                        icon=ft.Icons.FOLDER_OPEN,
                        on_click=self._open_file,
                    ),
                    ft.Button(
                        content="Нова книга",
                        icon=ft.Icons.NOTE_ADD,
                        on_click=self._new_file,
                    ),
                    ft.Button(
                        content="Зберегти",
                        icon=ft.Icons.SAVE,
                        on_click=self._save_file,
                    ),
                    ft.Button(
                        content="Зберегти як",
                        icon=ft.Icons.SAVE_AS,
                        on_click=self._save_as,
                    ),
                    ft.Button(
                        content="Оновити поточний стан",
                        icon=ft.Icons.REFRESH,
                        on_click=self._rebuild_current,
                    ),
                    ft.Container(expand=True),
                    ft.Text(f"v{APP_VERSION}", color=ft.Colors.WHITE_70),
                ],
            ),
        )

        nav_controls = [
            ft.Text(
                "РОЗДІЛИ",
                size=11,
                weight=ft.FontWeight.BOLD,
                color=ft.Colors.BLUE_GREY_400,
            )
        ]
        for sheet in REQUIRED_SHEETS:
            button = ft.Button(
                content=sheet,
                data=sheet,
                on_click=self._select_sheet,
                width=205,
            )
            self.sheet_buttons[sheet] = button
            nav_controls.append(button)

        nav_controls.extend([
            ft.Divider(),
            ft.Button(
                content="Налаштування",
                icon=ft.Icons.SETTINGS,
                data=SETTINGS_VIEW,
                on_click=self._select_sheet,
                width=205,
            ),
        ])

        sidebar = ft.Container(
            width=230,
            bgcolor=ft.Colors.BLUE_GREY_50,
            padding=16,
            content=ft.Column(
                controls=nav_controls,
                spacing=8,
            ),
        )

        action_bar = ft.Row(
            controls=[
                self.search,
                self.summary_group_dropdown,
                self.btn_summary_refresh,
                self.btn_add,
                self.action_menu,
            ],
            spacing=8,
        )

        main_panel = ft.Container(
            expand=True,
            padding=16,
            content=ft.Column(
                expand=True,
                controls=[
                    ft.Row(
                        controls=[
                            ft.Column(
                                controls=[
                                    ft.Text(
                                        self.current_sheet,
                                        key="sheet_title",
                                        size=23,
                                        weight=ft.FontWeight.BOLD,
                                    ),
                                    self.file_label,
                                ],
                                spacing=2,
                                expand=True,
                            ),
                        ]
                    ),
                    action_bar,
                    ft.Divider(height=1),
                    self.table_host,
                    ft.Divider(height=1),
                    self.status,
                ],
            ),
        )
        self.sheet_title = main_panel.content.controls[0].controls[0].controls[0]

        body = ft.Row(
            expand=True,
            spacing=0,
            controls=[sidebar, ft.VerticalDivider(width=1), main_panel],
        )
        self.page.add(
            ft.Column(
                expand=True,
                spacing=0,
                controls=[toolbar, body],
            )
        )
        self._refresh_table()

    def _autoload_default_book(self):
        base = (
            Path(sys.executable).parent
            if getattr(sys, "frozen", False)
            else Path(__file__).resolve().parent.parent
        )
        for candidate in (base / "Oblik.xlsx", base / "data" / "Oblik.xlsx"):
            if candidate.exists():
                try:
                    self.model.load(candidate)
                    self._set_status(f"Автоматично відкрито: {candidate}")
                    self._refresh_table()
                except Exception as exc:
                    self._show_message("Помилка відкриття", str(exc))
                break

    def _set_status(self, text: str):
        self.status.value = text
        self.page.update()

    def _show_message(self, title: str, message: str):
        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(title),
            content=ft.Text(message),
            actions=[
                ft.TextButton("OK", on_click=lambda e: self.page.pop_dialog())
            ],
        )
        self.page.show_dialog(dialog)

    async def _open_file(self, e=None):
        try:
            files = await self.file_picker.pick_files(
                dialog_title="Відкрити книгу обліку",
                allow_multiple=False,
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx", "xlsm"],
            )
            if not files:
                return
            selected = files[0]
            if not selected.path:
                self._show_message(
                    "Файл недоступний",
                    "Не вдалося отримати локальний шлях до вибраного файлу.",
                )
                return
            self.model.load(Path(selected.path))
            self.selected_excel_row = None
            self._set_status(f"Відкрито: {selected.path}")
            self._refresh_table()
        except Exception as exc:
            self._show_message("Помилка відкриття", str(exc))

    async def _new_file(self, e=None):
        try:
            path = await self.file_picker.save_file(
                dialog_title="Створити книгу",
                file_name="Oblik.xlsx",
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx"],
            )
            if not path:
                return
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.model.create_new(Path(path))
            self.current_sheet = SHEET_MOVEMENT
            self.selected_excel_row = None
            self._set_status(f"Створено: {path}")
            self._refresh_table()
        except Exception as exc:
            self._show_message("Помилка створення", str(exc))

    def _save_file(self, e=None):
        if self.model.wb is None:
            self._show_message("Збереження", "Спочатку відкрийте або створіть книгу.")
            return
        try:
            self.model.save()
            self._set_status(f"Збережено: {self.model.path}")
        except PermissionError:
            self._show_message(
                "Файл зайнятий",
                "Закрийте книгу в Excel і повторіть збереження.",
            )
        except Exception as exc:
            self._show_message("Помилка збереження", str(exc))

    async def _save_as(self, e=None):
        if self.model.wb is None:
            self._show_message("Збереження", "Спочатку відкрийте або створіть книгу.")
            return
        try:
            path = await self.file_picker.save_file(
                dialog_title="Зберегти копію",
                file_name="Oblik_copy.xlsx",
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx"],
            )
            if not path:
                return
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.model.save(Path(path))
            self._set_status(f"Збережено копію: {path}")
            self._refresh_table()
        except Exception as exc:
            self._show_message("Помилка збереження", str(exc))

    def _select_sheet(self, e):
        sheet = e.control.data
        if not sheet:
            return
        self.current_sheet = sheet
        self.selected_excel_row = None
        self.search.value = ""
        self.sheet_title.value = sheet
        self._refresh_table()

    def _search_changed(self, e=None):
        self.selected_excel_row = None
        self._refresh_table()

    def _settings_field(self, key: str, label: str, hint: str = "") -> ft.TextField:
        field = ft.TextField(
            label=label,
            hint_text=hint or None,
            value=self.settings.get(key, ""),
        )
        self.settings_fields[key] = field
        return field

    def _settings_card(self, title: str, subtitle: str, controls: list[ft.Control]) -> ft.Container:
        return ft.Container(
            padding=18,
            border=ft.Border.all(1, ft.Colors.BLUE_GREY_100),
            border_radius=12,
            bgcolor=ft.Colors.WHITE,
            content=ft.Column(
                controls=[
                    ft.Text(title, size=18, weight=ft.FontWeight.BOLD),
                    ft.Text(subtitle, size=12, color=ft.Colors.BLUE_GREY_600),
                    ft.Divider(),
                    *controls,
                ],
                spacing=10,
            ),
        )

    def _render_settings(self):
        self.table_host.controls.clear()
        self.search.visible = False
        self.btn_add.visible = False
        self.btn_edit.visible = False
        self.btn_delete.visible = False
        self.action_menu.visible = False
        self.summary_group_dropdown.visible = False
        self.btn_summary_refresh.visible = False
        self.sheet_title.value = SETTINGS_VIEW
        self.file_label.value = f"Файл налаштувань: {self.settings_store.path.name}"
        self.settings_fields = {}

        unit_card = self._settings_card(
            "Військова частина",
            "Основні реквізити, які надалі можна використовувати у відомостях і звітах.",
            [
                self._settings_field("unit_number", "Номер військової частини", "Наприклад: А0000"),
                self._settings_field("unit_name", "Найменування військової частини"),
            ],
        )

        commander_card = self._settings_card(
            "Командир військової частини",
            "Дані для автоматичного підставлення у службові документи.",
            [
                self._settings_field("commander_position", "Посада"),
                self._settings_field("commander_rank", "Військове звання"),
                self._settings_field("commander_name", "Ім’я та Прізвище"),
            ],
        )

        service_card = self._settings_card(
            "Начальник служби / відповідальна особа",
            "Реквізити служби, яка веде облік.",
            [
                self._settings_field("service_name", "Назва служби", "Наприклад: служба озброєння"),
                self._settings_field("service_chief_position", "Посада"),
                self._settings_field("service_chief_rank", "Військове звання"),
                self._settings_field("service_chief_name", "Ім’я та Прізвище"),
            ],
        )

        finance_card = self._settings_card(
            "Фінансово-економічна служба",
            "Дані начальника ФЕС для майбутніх документів і погоджень.",
            [
                self._settings_field("finance_chief_position", "Посада"),
                self._settings_field("finance_chief_rank", "Військове звання"),
                self._settings_field("finance_chief_name", "Ім’я та Прізвище"),
            ],
        )

        coefficients_card = self._settings_card(
            "Коефіцієнти індексації",
            "Постійні коефіцієнти за роками, як у старому аркуші Setting. Вони лише зберігаються в налаштуваннях і поки автоматично нічого не перераховують.",
            [
                self._settings_field("index_coefficient_2023", "Коефіцієнт 2023"),
                self._settings_field("index_coefficient_2024", "Коефіцієнт 2024"),
                self._settings_field("index_coefficient_2025", "Коефіцієнт 2025"),
            ],
        )

        inventory_prefix_field = self._settings_field(
            "inventory_prefix", "Префікс інвентарного номера", "Наприклад: ОВТ-"
        )
        inventory_suffix_field = self._settings_field(
            "inventory_suffix", "Суфікс інвентарного номера", "Наприклад: /26"
        )
        inventory_next_field = self._settings_field(
            "inventory_next_number", "Наступний порядковий номер", "Наприклад: 1"
        )
        inventory_digits_field = self._settings_field(
            "inventory_digits", "Кількість цифр у номері", "Наприклад: 6"
        )
        inventory_preview = ft.Text(
            "",
            size=14,
            weight=ft.FontWeight.W_600,
            color=ft.Colors.BLUE_GREY_800,
        )

        def refresh_inventory_preview(e=None):
            number = safe_positive_int(inventory_next_field.value, 1)
            inventory_preview.value = (
                "Приклад: "
                + format_inventory_number(
                    inventory_prefix_field.value,
                    number,
                    inventory_digits_field.value,
                    inventory_suffix_field.value,
                )
            )
            self.page.update()

        inventory_prefix_field.on_change = refresh_inventory_preview
        inventory_suffix_field.on_change = refresh_inventory_preview
        inventory_next_field.on_change = refresh_inventory_preview
        inventory_digits_field.on_change = refresh_inventory_preview
        refresh_inventory_preview()

        inventory_generator_card = self._settings_card(
            "Генератор інвентарних номерів",
            "Формат наступного інвентарного номера. При генерації програма автоматично пропускає номери, які вже є в «Рух майна».",
            [
                inventory_prefix_field,
                inventory_suffix_field,
                inventory_next_field,
                inventory_digits_field,
                inventory_preview,
                ft.Text(
                    "Наприклад: префікс ОВТ-, номер 25, 6 цифр → ОВТ-000025.",
                    size=12,
                    color=ft.Colors.BLUE_GREY_600,
                ),
            ],
        )

        numbering_card = self._settings_card(
            "Нумерація документів",
            "Загальні параметри, які можна використовувати при автоматичному формуванні документів.",
            [
                self._settings_field("document_prefix", "Префікс номера", "Наприклад: ЗВ-"),
                self._settings_field("document_start_number", "Початковий номер", "Наприклад: 1"),
            ],
        )

        self.commission_member_entries = []
        commission_members_host = ft.Column(spacing=10)
        commission_count = ft.Text(
            "Членів комісії: 0",
            size=12,
            color=ft.Colors.BLUE_GREY_600,
        )

        def refresh_member_titles():
            for index, entry in enumerate(self.commission_member_entries, start=1):
                entry["title"].value = f"Член комісії №{index}"
            commission_count.value = f"Членів комісії: {len(self.commission_member_entries)}"

        def add_commission_member(e=None, member=None):
            member = member if isinstance(member, dict) else {}
            position = ft.TextField(
                label="Посада",
                value=str(member.get("position", "") or ""),
                expand=2,
            )
            rank = ft.TextField(
                label="Військове звання",
                value=str(member.get("rank", "") or ""),
                expand=1,
            )
            name = ft.TextField(
                label="Ім’я та Прізвище",
                value=str(member.get("name", "") or ""),
                expand=2,
            )
            title = ft.Text(
                "Член комісії",
                weight=ft.FontWeight.W_600,
                color=ft.Colors.BLUE_GREY_700,
            )
            entry: dict[str, Any] = {
                "position": position,
                "rank": rank,
                "name": name,
                "title": title,
            }

            def remove_member(evt=None):
                if entry in self.commission_member_entries:
                    self.commission_member_entries.remove(entry)
                if entry["container"] in commission_members_host.controls:
                    commission_members_host.controls.remove(entry["container"])
                refresh_member_titles()
                self.page.update()

            container = ft.Container(
                padding=12,
                bgcolor=ft.Colors.BLUE_GREY_50,
                border_radius=10,
                content=ft.Column(
                    controls=[
                        ft.Row(
                            controls=[
                                title,
                                ft.Container(expand=True),
                                ft.IconButton(
                                    icon=ft.Icons.DELETE_OUTLINE,
                                    tooltip="Видалити члена комісії",
                                    on_click=remove_member,
                                ),
                            ]
                        ),
                        ft.Row(
                            controls=[position, rank, name],
                            spacing=8,
                        ),
                    ],
                    spacing=8,
                ),
            )
            entry["container"] = container
            self.commission_member_entries.append(entry)
            commission_members_host.controls.append(container)
            refresh_member_titles()
            if e is not None:
                self.page.update()

        existing_members = self.settings.get("commission_members", [])
        if isinstance(existing_members, list):
            for member in existing_members:
                add_commission_member(member=member)

        commission_card = self._settings_card(
            "Комісія",
            "Голова комісії та довільна кількість членів. Склад комісії можна змінювати без обмеження кількості.",
            [
                ft.Text(
                    "Голова комісії",
                    weight=ft.FontWeight.BOLD,
                    color=ft.Colors.BLUE_GREY_800,
                ),
                self._settings_field("commission_chair_position", "Посада голови комісії"),
                self._settings_field("commission_chair_rank", "Військове звання голови"),
                self._settings_field("commission_chair_name", "Ім’я та Прізвище голови комісії"),
                ft.Divider(),
                ft.Row(
                    controls=[
                        ft.Column(
                            controls=[
                                ft.Text(
                                    "Члени комісії",
                                    weight=ft.FontWeight.BOLD,
                                ),
                                commission_count,
                            ],
                            spacing=1,
                        ),
                        ft.Container(expand=True),
                        ft.Button(
                            content="+ Додати члена комісії",
                            icon=ft.Icons.PERSON_ADD,
                            on_click=add_commission_member,
                        ),
                    ]
                ),
                commission_members_host,
            ],
        )

        def save_settings(e=None):
            values = {
                key: field.value or ""
                for key, field in self.settings_fields.items()
            }
            values["commission_members"] = [
                {
                    "position": entry["position"].value or "",
                    "rank": entry["rank"].value or "",
                    "name": entry["name"].value or "",
                }
                for entry in self.commission_member_entries
                if any([
                    (entry["position"].value or "").strip(),
                    (entry["rank"].value or "").strip(),
                    (entry["name"].value or "").strip(),
                ])
            ]
            try:
                path = self.settings_store.save(values)
                self.settings = self.settings_store.load()
                self.status.value = f"Налаштування збережено: {path}"
                self._show_message(
                    "Налаштування збережено",
                    "Реквізити збережені локально і будуть доступні після наступного запуску програми.",
                )
            except PermissionError:
                self._show_message(
                    "Немає доступу до запису",
                    "Не вдалося зберегти oblik_settings.json поруч із програмою. Перемістіть Oblik.exe у папку, де дозволений запис.",
                )
            except Exception as exc:
                self._show_message("Помилка налаштувань", str(exc))

        self.table_host.controls.extend([
            ft.Container(
                expand=True,
                content=ft.Column(
                    expand=True,
                    scroll=ft.ScrollMode.AUTO,
                    controls=[
                        ft.Container(
                            width=900,
                            content=ft.Column(
                                controls=[
                        ft.Text(
                            "Ці реквізити зберігаються локально на цьому комп'ютері та не записуються в Excel автоматично.",
                            color=ft.Colors.BLUE_GREY_700,
                        ),
                        unit_card,
                        commander_card,
                        service_card,
                        finance_card,
                        coefficients_card,
                        commission_card,
                        inventory_generator_card,
                        numbering_card,
                                    ft.Row(
                                        alignment=ft.MainAxisAlignment.END,
                                        controls=[
                                            ft.Button(
                                                content="Зберегти налаштування",
                                                icon=ft.Icons.SAVE,
                                                on_click=save_settings,
                                            ),
                                        ],
                                    ),
                                ],
                                spacing=14,
                            ),
                        )
                    ],
                ),
            )
        ])
        self.status.value = f"Налаштування: {self.settings_store.path}"
        self.page.update()

    def _update_edit_permissions(self):
        editable = self.current_sheet == SHEET_MOVEMENT and self.model.wb is not None
        self.btn_add.disabled = not editable
        selected = editable and self.selected_excel_row is not None
        self.btn_edit.disabled = not selected
        self.btn_delete.disabled = not selected
        self.action_menu.disabled = not selected
        self.action_menu.opacity = 1.0 if selected else 0.45

    def _summary_group_changed(self, e=None):
        selected = (
            (e.control.value if e is not None and e.control is not None else None)
            or self.summary_group_dropdown.value
            or self.summary_group_mode
        )
        if selected in SUMMARY_GROUP_FIELDS:
            self.summary_group_mode = selected
            self.summary_group_dropdown.value = selected
        self._render_summary(sync_sheet=True)

    def _refresh_summary(self, e=None):
        if self.model.wb is None:
            self._show_message("Зведений", "Спочатку відкрийте або створіть книгу.")
            return
        try:
            self.model.rebuild_current_state()
            self._render_summary(sync_sheet=True)
            self._set_status(
                f"Зведений оновлено: групування «{self.summary_group_mode}»."
            )
        except Exception as exc:
            self._show_message("Помилка зведення", str(exc))

    def _render_summary(self, sync_sheet: bool = True):
        self.table_host.controls.clear()
        self.sheet_title.value = SHEET_SUMMARY
        self.file_label.value = (
            str(self.model.path) if self.model.path else "Файл не відкрито"
        )

        self.search.visible = True
        self.btn_add.visible = False
        self.btn_edit.visible = False
        self.btn_delete.visible = False
        self.action_menu.visible = False
        self.summary_group_dropdown.visible = True
        self.btn_summary_refresh.visible = True
        self.summary_group_dropdown.value = self.summary_group_mode

        if self.model.wb is None:
            self.table_host.controls.append(
                ft.Container(
                    padding=30,
                    content=ft.Text(
                        "Відкрийте Excel-файл або створіть нову книгу.",
                        color=ft.Colors.BLUE_GREY_600,
                    ),
                )
            )
            self.page.update()
            return

        group_header = SUMMARY_GROUP_FIELDS[self.summary_group_mode]
        summary = (
            self.model.rebuild_summary(group_header)
            if sync_sheet
            else self.model.build_summary_dataframe(group_header)
        )

        query = norm(self.search.value)
        if query and not summary.empty:
            mask = summary.apply(
                lambda row: any(
                    query in norm(display_value(value))
                    for value in row.values
                ),
                axis=1,
            )
            summary = summary[mask]

        visible_headers = [h for h in SUMMARY_OUTPUT_HEADERS if h != "№ з/п"]
        columns = [
            ft.DataColumn(
                label=ft.Container(
                    width=190 if header != "Значення" else 320,
                    padding=4,
                    content=ft.Text(
                        header,
                        size=12,
                        weight=ft.FontWeight.BOLD,
                    ),
                )
            )
            for header in visible_headers
        ]

        rows = []
        for _, record in summary.iterrows():
            cells = []
            for header in visible_headers:
                value = record.get(header)
                if header == "Загальна сума":
                    cell_text = format_decimal(numeric_value(value), 2)
                elif header in (
                    "Загальна кількість",
                    "Справні",
                    "Несправні",
                    "В ремонті",
                    "Знищені",
                    "Втрачені",
                    "Списані",
                ):
                    number = numeric_value(value)
                    if number is None:
                        cell_text = ""
                    elif float(number).is_integer():
                        cell_text = str(int(number))
                    else:
                        cell_text = format_decimal(number, 2)
                else:
                    cell_text = display_value(value)
                cells.append(
                    ft.DataCell(
                        ft.Container(
                            width=190 if header != "Значення" else 320,
                            padding=4,
                            content=ft.Text(cell_text, size=12, max_lines=3),
                        )
                    )
                )
            rows.append(ft.DataRow(cells=cells))

        table = fdt.DataTable2(
            columns=columns,
            rows=rows,
            expand=True,
            heading_row_color=ft.Colors.BLUE_50,
            fixed_top_rows=1,
            fixed_left_columns=1,
            fixed_columns_color=ft.Colors.BLUE_GREY_50,
            fixed_corner_color=ft.Colors.BLUE_100,
            visible_horizontal_scroll_bar=True,
            visible_vertical_scroll_bar=True,
            min_width=max(1250, len(visible_headers) * 205),
            heading_row_height=58,
            data_row_height=58,
            column_spacing=8,
            horizontal_margin=8,
        )

        self.table_host.controls.extend([
            ft.Container(
                padding=ft.Padding.only(bottom=8),
                content=ft.Text(
                    "Джерело: «Поточний стан». Зміна режиму не змінює дані — "
                    "лише спосіб їх групування.",
                    size=12,
                    color=ft.Colors.BLUE_GREY_600,
                ),
            ),
            ft.Container(expand=True, content=table),
        ])

        self.status.value = (
            f"Зведений: {len(summary)} груп | "
            f"режим «{self.summary_group_mode}» | {self.model.path or ''}"
        )
        self.page.update()

    def _refresh_table(self):
        if self.current_sheet == SETTINGS_VIEW:
            self._render_settings()
            return
        if self.current_sheet == SHEET_SUMMARY:
            self._render_summary(sync_sheet=True)
            return

        self.search.visible = True
        self.btn_add.visible = self.current_sheet == SHEET_MOVEMENT
        self.btn_edit.visible = False
        self.btn_delete.visible = False
        self.action_menu.visible = self.current_sheet == SHEET_MOVEMENT
        self.summary_group_dropdown.visible = False
        self.btn_summary_refresh.visible = False
        self.table_host.controls.clear()
        self.sheet_title.value = self.current_sheet
        self.file_label.value = (
            str(self.model.path) if self.model.path else "Файл не відкрито"
        )

        if self.model.wb is None or self.current_sheet not in self.model.wb.sheetnames:
            self.table_host.controls.append(
                ft.Container(
                    padding=30,
                    content=ft.Text(
                        "Відкрийте Excel-файл або створіть нову книгу.",
                        color=ft.Colors.BLUE_GREY_600,
                    ),
                )
            )
            self._update_edit_permissions()
            self.page.update()
            return

        df = self.model.dataframe(self.current_sheet)
        headers = self.model.headers(self.current_sheet)
        visible_headers = [h for h in headers if h != "№ з/п"]
        if TRANSACTION_ID_HEADER in visible_headers:
            visible_headers.remove(TRANSACTION_ID_HEADER)
            visible_headers.insert(0, TRANSACTION_ID_HEADER)
        if OPERATION_TYPE_HEADER in visible_headers:
            visible_headers.remove(OPERATION_TYPE_HEADER)
            insert_at = 1 if TRANSACTION_ID_HEADER in visible_headers else 0
            visible_headers.insert(insert_at, OPERATION_TYPE_HEADER)
        query = norm(self.search.value)

        duplicate_map = (
            self.model.duplicate_map()
            if self.current_sheet == SHEET_MOVEMENT
            else {}
        )

        data_rows = []
        matched = 0
        for _, row in df.iterrows():
            if query:
                searchable = []
                for header in visible_headers:
                    if header == "Сума":
                        value = calculate_total(row.get("Ціна"), row.get("Кількість"))
                    else:
                        value = row.get(header)
                    searchable.append(norm(display_value(value)))
                if not any(query in value for value in searchable):
                    continue

            matched += 1
            if len(data_rows) >= MAX_TABLE_ROWS:
                continue

            excel_row = int(row["_excel_row"])
            dup = duplicate_map.get(excel_row, DuplicateInfo())
            text_color = None
            if dup.score >= 5:
                text_color = ft.Colors.RED_700
            elif dup.score == 4:
                text_color = ft.Colors.ORANGE_800
            elif dup.score == 3:
                text_color = ft.Colors.AMBER_900

            cells = []
            for header in visible_headers:
                if header == "Сума":
                    value = calculate_total(row.get("Ціна"), row.get("Кількість"))
                    cell_text = "" if value is None else format_decimal(value, 2)
                else:
                    cell_text = display_value(row.get(header))

                cells.append(
                    ft.DataCell(
                        ft.Container(
                            width=165,
                            padding=4,
                            content=ft.Text(
                                cell_text,
                                size=12,
                                color=text_color,
                                max_lines=3,
                            ),
                        )
                    )
                )

            data_rows.append(
                ft.DataRow(
                    data=excel_row,
                    selected=excel_row == self.selected_excel_row,
                    on_select_change=self._row_selected,
                    cells=cells,
                )
            )

        columns = [
            ft.DataColumn(
                label=ft.Container(
                    width=165,
                    padding=4,
                    content=ft.Text(
                        header.replace("\n", " "),
                        size=12,
                        weight=ft.FontWeight.BOLD,
                    ),
                )
            )
            for header in visible_headers
        ]

        table = fdt.DataTable2(
            columns=columns,
            rows=data_rows,
            expand=True,
            show_checkbox_column=False,
            heading_row_color=ft.Colors.BLUE_50,
            fixed_top_rows=1,
            fixed_left_columns=1 if TRANSACTION_ID_HEADER in visible_headers else 0,
            fixed_columns_color=ft.Colors.BLUE_GREY_50,
            fixed_corner_color=ft.Colors.BLUE_100,
            visible_horizontal_scroll_bar=True,
            visible_vertical_scroll_bar=True,
            min_width=max(1100, len(visible_headers) * 181),
            heading_row_height=58,
            data_row_height=64,
            column_spacing=8,
            horizontal_margin=8,
        )

        # DataTable2 тримає шапку зверху та ID транзакції зліва.
        # Прокручується лише тіло таблиці, а обидва scrollbars завжди доступні.
        self.table_host.controls.append(
            ft.Container(
                expand=True,
                content=table,
            )
        )

        shown = min(matched, MAX_TABLE_ROWS)
        suffix = (
            f" | показано перші {MAX_TABLE_ROWS}"
            if matched > MAX_TABLE_ROWS
            else ""
        )
        self.status.value = (
            f"{self.current_sheet}: {matched} записів{suffix}"
            f" | {self.model.path or ''}"
        )
        self._update_edit_permissions()
        self.page.update()

    def _row_selected(self, e):
        excel_row = int(e.control.data)
        is_selected = bool(e.data)
        self.selected_excel_row = excel_row if is_selected else None
        self._refresh_table()

    def _selected_record_values(self) -> Optional[dict[str, Any]]:
        if (
            self.model.wb is None
            or self.current_sheet != SHEET_MOVEMENT
            or self.selected_excel_row is None
        ):
            self._show_message("Дія", "Спочатку виберіть запис у таблиці.")
            return None
        ws = self.model.wb[SHEET_MOVEMENT]
        headers = self.model.headers(SHEET_MOVEMENT)
        return {
            header: ws.cell(self.selected_excel_row, i + 1).value
            for i, header in enumerate(headers)
        }

    @staticmethod
    def _append_action_note(existing: Any, line: str) -> str:
        current = display_value(existing).strip()
        return f"{current}\n{line}".strip() if current else line

    def _create_action_transaction(
        self,
        base: dict[str, Any],
        operation_type: str,
        updates: dict[str, Any],
        note_line: str = "",
    ) -> int:
        data = dict(base)
        source_id = display_value(base.get(TRANSACTION_ID_HEADER)).strip()
        data[TRANSACTION_ID_HEADER] = None
        data[OPERATION_TYPE_HEADER] = operation_type
        data.update(updates)
        if note_line:
            data["Примітка"] = self._append_action_note(
                data.get("Примітка"),
                note_line,
            )
        reason = (
            f"{operation_type}; базова транзакція {source_id}"
            if source_id
            else operation_type
        )
        new_row = self.model.append_record(SHEET_MOVEMENT, data, reason)
        try:
            self.model.rebuild_current_state()
        except Exception:
            # Нова транзакція все одно зберігається; поточний стан можна
            # перебудувати окремою кнопкою.
            pass
        self.selected_excel_row = new_row
        self._refresh_table()
        self._set_status(
            f"Створено нову транзакцію: {operation_type}. "
            "Не забудьте зберегти книгу."
        )
        return new_row

    def _open_move_dialog(self, e=None):
        base = self._selected_record_values()
        if not base:
            return

        from_brigade = display_value(base.get(MAIN_HEADERS[19]))
        from_subunit = display_value(base.get(MAIN_HEADERS[20]))
        to_brigade = ft.TextField(label="Куди — підрозділ бригади")
        to_subunit = ft.TextField(label="Куди — підрозділ батальйону/дивізіону")
        doc_no = ft.TextField(label="Номер документа")
        doc_date = ft.TextField(label="Дата документа", hint_text="дд.мм.рррр")
        note = ft.TextField(label="Примітка", multiline=True, min_lines=2, max_lines=4)

        def save_move(evt=None):
            if not (to_brigade.value or "").strip() and not (to_subunit.value or "").strip():
                self._show_message(
                    "Переміщення",
                    "Вкажіть хоча б один підрозділ призначення.",
                )
                return
            updates = {
                MAIN_HEADERS[19]: parse_user_value(MAIN_HEADERS[19], to_brigade.value or ""),
                MAIN_HEADERS[20]: parse_user_value(MAIN_HEADERS[20], to_subunit.value or ""),
                MAIN_HEADERS[21]: parse_user_value(MAIN_HEADERS[21], doc_no.value or ""),
                MAIN_HEADERS[22]: parse_user_value(MAIN_HEADERS[22], doc_date.value or ""),
            }
            destination = " / ".join(
                value for value in [
                    (to_brigade.value or "").strip(),
                    (to_subunit.value or "").strip(),
                ]
                if value
            )
            extra = (note.value or "").strip()
            line = f"[Переміщення] → {destination}"
            if extra:
                line += f"; {extra}"
            self.page.pop_dialog()
            self._create_action_transaction(base, "Переміщення", updates, line)

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Перемістити майно"),
            content=ft.Container(
                width=700,
                content=ft.Column(
                    tight=True,
                    controls=[
                        ft.Text(
                            f"Транзакція: {display_value(base.get(TRANSACTION_ID_HEADER))}",
                            weight=ft.FontWeight.BOLD,
                        ),
                        ft.Text(
                            "Звідки: "
                            + (" / ".join(x for x in [from_brigade, from_subunit] if x) or "не вказано"),
                            color=ft.Colors.BLUE_GREY_700,
                        ),
                        to_brigade,
                        to_subunit,
                        ft.Row(controls=[doc_no, doc_date], spacing=8),
                        note,
                    ],
                ),
            ),
            actions=[
                ft.TextButton("Скасувати", on_click=lambda evt: self.page.pop_dialog()),
                ft.Button(content="Створити переміщення", icon=ft.Icons.SWAP_HORIZ, on_click=save_move),
            ],
        )
        self.page.show_dialog(dialog)

    def _open_writeoff_dialog(self, e=None):
        base = self._selected_record_values()
        if not base:
            return

        order_no = ft.TextField(label="Номер наказу на списання")
        order_date = ft.TextField(label="Дата наказу на списання", hint_text="дд.мм.рррр")
        act_no = ft.TextField(label="Номер Єдиного акту списання")
        act_date = ft.TextField(label="Дата Єдиного акту списання", hint_text="дд.мм.рррр")
        note = ft.TextField(label="Примітка", multiline=True, min_lines=2, max_lines=4)

        def save_writeoff(evt=None):
            parsed_order_date = parse_user_value(MAIN_HEADERS[30], order_date.value or "")
            parsed_act_date = parse_user_value(MAIN_HEADERS[32], act_date.value or "")
            status_date = parsed_act_date or parsed_order_date
            updates = {
                MAIN_HEADERS[23]: "Списаний",
                MAIN_HEADERS[24]: status_date,
                MAIN_HEADERS[29]: parse_user_value(MAIN_HEADERS[29], order_no.value or ""),
                MAIN_HEADERS[30]: parsed_order_date,
                MAIN_HEADERS[31]: parse_user_value(MAIN_HEADERS[31], act_no.value or ""),
                MAIN_HEADERS[32]: parsed_act_date,
            }
            details = []
            if (order_no.value or "").strip():
                details.append(f"наказ №{order_no.value.strip()}")
            if (act_no.value or "").strip():
                details.append(f"акт №{act_no.value.strip()}")
            extra = (note.value or "").strip()
            line = "[Списання]"
            if details:
                line += " " + ", ".join(details)
            if extra:
                line += f"; {extra}"
            self.page.pop_dialog()
            self._create_action_transaction(base, "Списання", updates, line)

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Списати майно"),
            content=ft.Container(
                width=720,
                content=ft.Column(
                    tight=True,
                    controls=[
                        ft.Text(
                            f"Транзакція: {display_value(base.get(TRANSACTION_ID_HEADER))}",
                            weight=ft.FontWeight.BOLD,
                        ),
                        ft.Row(controls=[order_no, order_date], spacing=8),
                        ft.Row(controls=[act_no, act_date], spacing=8),
                        note,
                    ],
                ),
            ),
            actions=[
                ft.TextButton("Скасувати", on_click=lambda evt: self.page.pop_dialog()),
                ft.Button(content="Створити списання", icon=ft.Icons.DELETE_SWEEP, on_click=save_writeoff),
            ],
        )
        self.page.show_dialog(dialog)

    def _open_status_dialog(self, e=None):
        base = self._selected_record_values()
        if not base:
            return

        states = [
            "Справний",
            "Несправний",
            "Переданий в ремонт",
            "Знищений",
            "Втрачений",
            "Списаний",
        ]
        current = display_value(base.get(MAIN_HEADERS[23]))
        status = ft.Dropdown(
            label="Новий стан",
            value=current if current in states else None,
            options=[ft.DropdownOption(key=value, text=value) for value in states],
        )
        status_date = ft.TextField(label="Дата зміни стану", hint_text="дд.мм.рррр")
        note = ft.TextField(label="Примітка", multiline=True, min_lines=2, max_lines=4)

        def save_status(evt=None):
            chosen = status.value or status.text or ""
            if not chosen:
                self._show_message("Зміна стану", "Оберіть новий стан.")
                return
            updates = {
                MAIN_HEADERS[23]: chosen,
                MAIN_HEADERS[24]: parse_user_value(MAIN_HEADERS[24], status_date.value or ""),
            }
            extra = (note.value or "").strip()
            line = f"[Зміна стану] {chosen}"
            if extra:
                line += f"; {extra}"
            self.page.pop_dialog()
            self._create_action_transaction(base, "Зміна стану", updates, line)

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Змінити стан майна"),
            content=ft.Container(
                width=620,
                content=ft.Column(
                    tight=True,
                    controls=[
                        ft.Text(
                            f"Транзакція: {display_value(base.get(TRANSACTION_ID_HEADER))}",
                            weight=ft.FontWeight.BOLD,
                        ),
                        status,
                        status_date,
                        note,
                    ],
                ),
            ),
            actions=[
                ft.TextButton("Скасувати", on_click=lambda evt: self.page.pop_dialog()),
                ft.Button(content="Створити транзакцію", icon=ft.Icons.HEALTH_AND_SAFETY, on_click=save_status),
            ],
        )
        self.page.show_dialog(dialog)

    def _open_stock_quantity_dialog(self, operation_type: str, sign: int):
        base = self._selected_record_values()
        if not base:
            return
        if "запас" not in norm(base.get(MAIN_HEADERS[8])):
            self._show_message(
                operation_type,
                "Ця дія доступна для записів із типом майна «Запас».",
            )
            return

        quantity = ft.TextField(label="Кількість")
        destination = ft.TextField(
            label="Кому / підрозділ",
            hint_text="Для повернення можна вказати склад або підрозділ",
        )
        doc_no = ft.TextField(label="Номер документа")
        doc_date = ft.TextField(label="Дата документа", hint_text="дд.мм.рррр")
        note = ft.TextField(label="Примітка", multiline=True, min_lines=2, max_lines=4)

        def save_stock(evt=None):
            qty = numeric_value(quantity.value)
            if qty is None or qty <= 0:
                self._show_message(operation_type, "Вкажіть кількість більше нуля.")
                return
            signed_qty = abs(qty) * sign
            updates = {
                MAIN_HEADERS[16]: signed_qty,
                MAIN_HEADERS[20]: parse_user_value(MAIN_HEADERS[20], destination.value or ""),
                MAIN_HEADERS[21]: parse_user_value(MAIN_HEADERS[21], doc_no.value or ""),
                MAIN_HEADERS[22]: parse_user_value(MAIN_HEADERS[22], doc_date.value or ""),
            }
            extra = (note.value or "").strip()
            line = f"[{operation_type}] {abs(qty):g}"
            if (destination.value or "").strip():
                line += f" → {destination.value.strip()}"
            if extra:
                line += f"; {extra}"
            self.page.pop_dialog()
            self._create_action_transaction(base, operation_type, updates, line)

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(operation_type),
            content=ft.Container(
                width=650,
                content=ft.Column(
                    tight=True,
                    controls=[
                        ft.Text(
                            "Кількісний залишок запасів ще не агрегується автоматично; "
                            "ця дія вже створює окрему транзакцію руху.",
                            color=ft.Colors.ORANGE_800,
                        ),
                        quantity,
                        destination,
                        ft.Row(controls=[doc_no, doc_date], spacing=8),
                        note,
                    ],
                ),
            ),
            actions=[
                ft.TextButton("Скасувати", on_click=lambda evt: self.page.pop_dialog()),
                ft.Button(content="Створити транзакцію", icon=ft.Icons.SAVE, on_click=save_stock),
            ],
        )
        self.page.show_dialog(dialog)

    def _open_issue_stock_dialog(self, e=None):
        self._open_stock_quantity_dialog("Видача", -1)

    def _open_return_stock_dialog(self, e=None):
        self._open_stock_quantity_dialog("Повернення", 1)

    def _show_selected_history(self, e=None):
        base = self._selected_record_values()
        if not base:
            return

        df = self.model.dataframe(SHEET_MOVEMENT)
        inv = norm(base.get(MAIN_HEADERS[9]))
        serial = norm(base.get(MAIN_HEADERS[18]))
        code = norm(base.get(MAIN_HEADERS[7]))
        name = norm(base.get(MAIN_HEADERS[12]) or base.get(MAIN_HEADERS[11]))

        def same_asset(row) -> bool:
            if inv:
                return norm(row.get(MAIN_HEADERS[9])) == inv
            if serial:
                return (
                    norm(row.get(MAIN_HEADERS[18])) == serial
                    and norm(row.get(MAIN_HEADERS[12]) or row.get(MAIN_HEADERS[11])) == name
                )
            return (
                bool(code or name)
                and (not code or norm(row.get(MAIN_HEADERS[7])) == code)
                and (not name or norm(row.get(MAIN_HEADERS[12]) or row.get(MAIN_HEADERS[11])) == name)
            )

        records = [
            row for _, row in df.iterrows()
            if same_asset(row)
        ]

        def event_date(row) -> str:
            for header in (
                MAIN_HEADERS[32],
                MAIN_HEADERS[30],
                MAIN_HEADERS[24],
                MAIN_HEADERS[22],
                MAIN_HEADERS[4],
            ):
                value = row.get(header)
                if not is_blank(value):
                    return display_value(value)
            return ""

        cards = []
        for row in reversed(records):
            location = " / ".join(
                value for value in [
                    display_value(row.get(MAIN_HEADERS[19])).strip(),
                    display_value(row.get(MAIN_HEADERS[20])).strip(),
                ]
                if value
            )
            document = (
                display_value(row.get(MAIN_HEADERS[31])).strip()
                or display_value(row.get(MAIN_HEADERS[29])).strip()
                or display_value(row.get(MAIN_HEADERS[21])).strip()
                or display_value(row.get(MAIN_HEADERS[3])).strip()
            )
            cards.append(
                ft.Container(
                    padding=12,
                    border=ft.Border.all(1, ft.Colors.BLUE_GREY_100),
                    border_radius=10,
                    content=ft.Column(
                        spacing=4,
                        controls=[
                            ft.Row(
                                controls=[
                                    ft.Text(
                                        display_value(row.get(TRANSACTION_ID_HEADER)),
                                        weight=ft.FontWeight.BOLD,
                                    ),
                                    ft.Text(
                                        display_value(row.get(OPERATION_TYPE_HEADER)) or "Запис",
                                        color=ft.Colors.BLUE_700,
                                    ),
                                    ft.Container(expand=True),
                                    ft.Text(event_date(row), color=ft.Colors.BLUE_GREY_600),
                                ]
                            ),
                            ft.Text(
                                f"Документ: {document or '—'} | "
                                f"Кількість: {display_value(row.get(MAIN_HEADERS[16])) or '—'}"
                            ),
                            ft.Text(
                                f"Місце: {location or '—'} | "
                                f"Стан: {display_value(row.get(MAIN_HEADERS[23])) or '—'}"
                            ),
                            ft.Text(
                                display_value(row.get("Примітка")),
                                color=ft.Colors.BLUE_GREY_700,
                            ) if display_value(row.get("Примітка")).strip() else ft.Container(),
                        ],
                    ),
                )
            )

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(
                "Історія: "
                + (
                    display_value(base.get(MAIN_HEADERS[9]))
                    or display_value(base.get(MAIN_HEADERS[18]))
                    or display_value(base.get(MAIN_HEADERS[12]))
                    or "майно"
                )
            ),
            content=ft.Container(
                width=900,
                height=560,
                content=ft.Column(
                    controls=cards or [ft.Text("Історію не знайдено.")],
                    scroll=ft.ScrollMode.AUTO,
                    spacing=8,
                ),
            ),
            actions=[
                ft.TextButton("Закрити", on_click=lambda evt: self.page.pop_dialog()),
            ],
        )
        self.page.show_dialog(dialog)

    def _display_headers_for_form(self, headers: list[str]) -> list[str]:
        result = [
            h for h in headers
            if h not in ("№ з/п", TRANSACTION_ID_HEADER, OPERATION_TYPE_HEADER)
        ]
        if "Ціна" in result and "Кількість" in result:
            result.remove("Кількість")
            result.insert(result.index("Ціна"), "Кількість")
        return result

    def _build_record_data(
        self,
        field_controls: dict[str, ft.Control],
        reverse_mode: bool,
        calculated_price: Optional[float],
        calculated_sum: Optional[float],
        price_input: ft.TextField,
        sum_input: ft.TextField,
    ) -> dict[str, Any]:
        out: dict[str, Any] = {}
        for header, control in field_controls.items():
            if header == "Ціна":
                out[header] = (
                    calculated_price
                    if reverse_mode
                    else parse_user_value(header, price_input.value or "")
                )
                continue
            if header == "Сума":
                out[header] = (
                    numeric_value(sum_input.value)
                    if reverse_mode
                    else calculated_sum
                )
                continue
            if isinstance(control, ft.Dropdown):
                # В editable Dropdown беремо фактично введений текст першим,
                # щоб ручне уточнення не перекривалося старим selected value.
                raw = control.text or control.value or ""
            elif isinstance(control, ft.TextField):
                raw = control.value or ""
            else:
                raw = ""
            out[header] = parse_user_value(header, raw)
        return out

    def _open_record_dialog(self, excel_row: Optional[int] = None):
        if self.model.wb is None or self.current_sheet != SHEET_MOVEMENT:
            return

        headers = self.model.headers(SHEET_MOVEMENT)
        values: dict[str, Any] = {}
        if excel_row is not None:
            ws = self.model.wb[SHEET_MOVEMENT]
            values = {
                header: ws.cell(excel_row, i + 1).value
                for i, header in enumerate(headers)
            }

        state = {
            "reverse": False,
            "calculated_price": numeric_value(values.get("Ціна")),
            "calculated_sum": None,
            "warnings_acknowledged": False,
            "generated_inventory": None,
            "generated_inventory_counter": None,
        }
        field_controls: dict[str, ft.Control] = {}

        qty_input = ft.TextField(
            value=display_value(values.get("Кількість")),
            on_change=lambda e: refresh_calculation(),
        )
        price_input = ft.TextField(
            value=display_value(values.get("Ціна")),
            on_change=lambda e: refresh_calculation(),
        )
        price_label = ft.Text(
            "—",
            size=16,
            weight=ft.FontWeight.W_600,
        )
        price_label_box = ft.Container(
            padding=ft.Padding.symmetric(horizontal=12, vertical=13),
            bgcolor=ft.Colors.BLUE_GREY_50,
            border_radius=8,
            content=price_label,
            visible=False,
        )

        sum_label = ft.Text(
            "—",
            size=16,
            weight=ft.FontWeight.W_600,
        )
        sum_label_box = ft.Container(
            padding=ft.Padding.symmetric(horizontal=12, vertical=13),
            bgcolor=ft.Colors.BLUE_GREY_50,
            border_radius=8,
            content=sum_label,
        )
        sum_input = ft.TextField(
            hint_text="Введіть загальну вартість",
            visible=False,
            on_change=lambda e: refresh_calculation(),
        )

        mode_button = ft.Button(
            content="⟳  Розрахунок від суми",
        )
        warning_text = ft.Text(
            "",
            color=ft.Colors.ORANGE_900,
            weight=ft.FontWeight.W_600,
            visible=False,
        )

        price_host = ft.Column(
            controls=[price_input, price_label_box],
            spacing=0,
        )
        sum_host = ft.Column(
            controls=[sum_label_box, sum_input],
            spacing=0,
        )

        def refresh_calculation():
            qty = numeric_value(qty_input.value)
            if state["reverse"]:
                state["calculated_sum"] = numeric_value(sum_input.value)
                state["calculated_price"] = calculate_unit_price(
                    sum_input.value, qty
                )
                price_label.value = format_decimal(
                    state["calculated_price"], 5
                )
            else:
                state["calculated_price"] = numeric_value(price_input.value)
                state["calculated_sum"] = calculate_total(
                    price_input.value, qty
                )
                sum_label.value = format_decimal(
                    state["calculated_sum"], 2
                )
            self.page.update()

        def toggle_mode(e=None):
            if state["reverse"]:
                if state["calculated_price"] is not None:
                    price_input.value = str(state["calculated_price"])
                state["reverse"] = False
            else:
                if not (sum_input.value or "").strip():
                    current_total = calculate_total(
                        price_input.value, qty_input.value
                    )
                    if current_total is not None:
                        sum_input.value = str(current_total)
                state["reverse"] = True

            price_input.visible = not state["reverse"]
            price_label_box.visible = state["reverse"]
            sum_label_box.visible = not state["reverse"]
            sum_input.visible = state["reverse"]
            mode_button.content = (
                "⟳  Розрахунок від суми: УВІМКНЕНО"
                if state["reverse"]
                else "⟳  Розрахунок від суми"
            )
            state["warnings_acknowledged"] = False
            warning_text.visible = False
            refresh_calculation()

        mode_button.on_click = toggle_mode

        inventory_input = ft.TextField(
            value=display_value(values.get(MAIN_HEADERS[9])),
            expand=True,
        )

        def generate_inventory(e=None):
            try:
                movement = self.model.dataframe(SHEET_MOVEMENT)
                used_values = (
                    movement[MAIN_HEADERS[9]].tolist()
                    if not movement.empty and MAIN_HEADERS[9] in movement.columns
                    else []
                )
                candidate, counter = next_inventory_number(self.settings, used_values)
                inventory_input.value = candidate
                state["generated_inventory"] = candidate
                state["generated_inventory_counter"] = counter
                warning_text.visible = False
                self.page.update()
            except Exception as exc:
                self._show_message("Генератор інвентарних номерів", str(exc))

        inventory_host = ft.Row(
            controls=[
                inventory_input,
                ft.Button(
                    content="Згенерувати",
                    icon=ft.Icons.AUTO_AWESOME,
                    on_click=generate_inventory,
                ),
            ],
            spacing=8,
        )

        form_rows = [
            ft.Container(
                padding=ft.Padding.only(bottom=8),
                content=ft.Text(
                    (
                        "ID транзакції: "
                        + str(values.get(TRANSACTION_ID_HEADER) or "")
                        if excel_row is not None
                        else "ID транзакції буде створено автоматично після збереження"
                    ),
                    size=12,
                    color=ft.Colors.BLUE_GREY_600,
                ),
            )
        ]
        display_headers = self._display_headers_for_form(headers)
        for header in display_headers:
            current = values.get(header)
            label = ft.Text(
                header.replace("\n", " "),
                size=12,
                weight=ft.FontWeight.W_600,
            )

            if header == "Кількість":
                control = qty_input
            elif header == "Ціна":
                control = price_host
            elif header == "Сума":
                control = sum_host
            elif header == MAIN_HEADERS[9]:
                control = inventory_input
                display_control = inventory_host
            elif header == MAIN_HEADERS[8]:
                current_text = display_value(current)
                options = ["", "Необоротний актив", "Запас"]
                control = ft.Dropdown(
                    editable=True,
                    text=current_text,
                    value=current_text if current_text in options and current_text else None,
                    options=[
                        ft.DropdownOption(key=o, text=o)
                        for o in options
                        if o
                    ],
                )
            elif header == MAIN_HEADERS[23]:
                current_text = display_value(current)
                options = [
                    "",
                    "Справний",
                    "Несправний",
                    "Переданий в ремонт",
                    "Знищений",
                    "Втрачений",
                    "Списаний",
                ]
                control = ft.Dropdown(
                    editable=True,
                    text=current_text,
                    value=current_text if current_text in options and current_text else None,
                    options=[
                        ft.DropdownOption(key=o, text=o)
                        for o in options
                        if o
                    ],
                )
            elif header == "Примітка":
                control = ft.TextField(
                    value=display_value(current),
                    multiline=True,
                    min_lines=2,
                    max_lines=4,
                )
            else:
                control = ft.TextField(
                    value=display_value(current),
                    hint_text=(
                        "дд.мм.рррр"
                        if "дата" in header.casefold()
                        else None
                    ),
                )

            field_controls[header] = control
            shown_control = locals().get("display_control", control)
            form_rows.append(
                ft.Row(
                    vertical_alignment=ft.CrossAxisAlignment.START,
                    controls=[
                        ft.Container(width=285, padding=8, content=label),
                        ft.Container(expand=True, content=shown_control),
                    ],
                )
            )
            if "display_control" in locals():
                del display_control

            if header == "Кількість":
                form_rows.append(
                    ft.Row(
                        controls=[
                            ft.Container(width=285),
                            ft.Container(expand=True, content=mode_button),
                        ]
                    )
                )

        reason_input = ft.TextField(
            label="Причина зміни / уточнення",
            hint_text="Наприклад: уточнення документів, помилка минулих років",
        )
        form_rows.append(ft.Divider())
        form_rows.append(reason_input)
        form_rows.append(warning_text)

        def commit_record(data: dict[str, Any]):
            try:
                if excel_row is None:
                    self.model.append_record(
                        SHEET_MOVEMENT,
                        data,
                        reason_input.value or "",
                    )
                    message = "Запис додано. Зміни внесено до 'Контроль змін'."
                else:
                    changes = self.model.update_record(
                        SHEET_MOVEMENT,
                        excel_row,
                        data,
                        reason_input.value or "",
                    )
                    message = (
                        f"Змінено полів: {len(changes)}. "
                        "Зміни внесено до 'Контроль змін'."
                    )
                generated_value = state.get("generated_inventory")
                generated_counter = state.get("generated_inventory_counter")
                if (
                    generated_value
                    and generated_counter is not None
                    and norm(data.get(MAIN_HEADERS[9])) == norm(generated_value)
                ):
                    self.settings["inventory_next_number"] = str(int(generated_counter) + 1)
                    self.settings_store.save(self.settings)
                    self.settings = self.settings_store.load()

                self.page.pop_dialog()
                self.selected_excel_row = None
                self._refresh_table()
                self._set_status(message)
            except Exception as exc:
                self._show_message("Помилка запису", str(exc))

        def save_record(e=None):
            refresh_calculation()
            if state["reverse"]:
                total = numeric_value(sum_input.value)
                qty = numeric_value(qty_input.value)
                if total is not None and (qty is None or qty == 0):
                    warning_text.value = (
                        "Для розрахунку ціни від суми вкажіть кількість, "
                        "більшу за нуль."
                    )
                    warning_text.visible = True
                    self.page.update()
                    return

            data = self._build_record_data(
                field_controls,
                state["reverse"],
                state["calculated_price"],
                state["calculated_sum"],
                price_input,
                sum_input,
            )

            warnings = []
            dup = self.model.score_candidate_duplicate(data, excel_row)
            if dup.score >= 3:
                if dup.score >= 5:
                    level = "ПОВНИЙ ДУБЛЬ 5/5"
                elif dup.score == 4:
                    level = "ДУЖЕ СХОЖИЙ ЗАПИС 4/5"
                else:
                    level = "МОЖЛИВИЙ ДУБЛЬ 3/5"
                warnings.append(
                    f"{level}. Схожі Excel-рядки: "
                    + ", ".join(map(str, dup.other_excel_rows))
                )

            asset_type = norm(data.get(MAIN_HEADERS[8]))
            if (
                "необорот" in asset_type
                and not norm(data.get(MAIN_HEADERS[9]))
            ):
                warnings.append(
                    "Для необоротного активу не вказаний інвентарний номер."
                )

            if warnings and not state["warnings_acknowledged"]:
                warning_text.value = (
                    "\n".join(warnings)
                    + "\n\nНатисніть «Зберегти все одно» ще раз, "
                    "якщо запис потрібно залишити."
                )
                warning_text.visible = True
                state["warnings_acknowledged"] = True
                save_button.content = "Зберегти все одно"
                self.page.update()
                return

            commit_record(data)

        save_button = ft.Button(
            content="Зберегти",
            icon=ft.Icons.SAVE,
            on_click=save_record,
        )
        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(
                "Редагувати запис" if excel_row is not None else "Додати запис"
            ),
            content=ft.Container(
                width=820,
                height=610,
                content=ft.Column(
                    controls=form_rows,
                    scroll=ft.ScrollMode.AUTO,
                    spacing=7,
                ),
            ),
            actions=[
                ft.TextButton(
                    "Скасувати",
                    on_click=lambda e: self.page.pop_dialog(),
                ),
                save_button,
            ],
        )
        refresh_calculation()
        self.page.show_dialog(dialog)

    def _add_record(self, e=None):
        self._open_record_dialog()

    def _edit_record(self, e=None):
        if self.selected_excel_row is None:
            self._show_message("Редагування", "Виберіть рядок.")
            return
        self._open_record_dialog(self.selected_excel_row)

    def _delete_record(self, e=None):
        if self.selected_excel_row is None or self.model.wb is None:
            return
        reason = ft.TextField(
            label="Причина видалення",
            hint_text="Наприклад: дубль, уточнення документів",
        )

        def confirm_delete(evt=None):
            try:
                row = self.selected_excel_row
                self.page.pop_dialog()
                if row is None:
                    return
                self.model.delete_record(
                    SHEET_MOVEMENT,
                    row,
                    reason.value or "",
                )
                self.selected_excel_row = None
                self._refresh_table()
                self._set_status(
                    "Запис видалено; інформація залишилась у 'Контроль змін'."
                )
            except Exception as exc:
                self._show_message("Помилка видалення", str(exc))

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Видалити запис?"),
            content=ft.Column(
                controls=[
                    ft.Text(
                        "Запис буде прибрано з таблиці, але інформація про "
                        "видалення залишиться у 'Контроль змін'."
                    ),
                    reason,
                ],
                tight=True,
            ),
            actions=[
                ft.TextButton(
                    "Скасувати",
                    on_click=lambda e: self.page.pop_dialog(),
                ),
                ft.Button(
                    content="Видалити",
                    icon=ft.Icons.DELETE,
                    on_click=confirm_delete,
                ),
            ],
        )
        self.page.show_dialog(dialog)

    def _rebuild_current(self, e=None):
        if self.model.wb is None:
            self._show_message(
                "Поточний стан",
                "Спочатку відкрийте або створіть книгу.",
            )
            return
        try:
            count, skipped = self.model.rebuild_current_state()
            self._refresh_table()
            self._show_message(
                "Поточний стан оновлено",
                f"Сформовано поштучних записів: {count}.\n"
                f"Записів без інвентарного/заводського номера, "
                f"які не згорнуті автоматично: {skipped}.\n\n"
                "Колонка 'Штатна потреба' не перераховувалась.",
            )
        except Exception as exc:
            self._show_message("Помилка розрахунку", str(exc))


def main(page: ft.Page):
    FletOblikApp(page)


if __name__ == "__main__":
    ft.run(main)
